import copy
import re
from collections import Counter
from io import BytesIO
from datetime import datetime, timezone

from bson import ObjectId
from pymongo import ASCENDING, DESCENDING
from rest_framework import status
from rest_framework.decorators import api_view, parser_classes
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response

from .mongo import get_db
from .utils import serialize_datetime, utcnow


RESUME_TEMPLATES = {
    "modern",
    "classic",
    "compact",
    "minimal",
}

SECTION_KEYS = (
    "experience",
    "education",
    "skills",
    "projects",
    "certifications",
    "languages",
)

STOP_WORDS = {
    "a", "an", "and", "are", "as", "at", "be", "been", "being",
    "by", "for", "from", "has", "have", "in", "into", "is", "it",
    "its", "of", "on", "or", "our", "that", "the", "their", "this",
    "to", "we", "will", "with", "you", "your", "they", "them", "who",
    "what", "when", "where", "which", "while", "work", "working",
    "role", "job", "position", "candidate", "candidates", "team",
    "teams", "company", "years", "year", "experience", "preferred",
    "required", "requirements", "responsibilities", "responsibility",
    "skills", "skill", "including", "such", "using", "use", "used",
}


def _owner_oid(request):
    return ObjectId(request.user.id)


def _oid(value):
    if isinstance(value, ObjectId):
        return value

    if not ObjectId.is_valid(str(value)):
        return None

    return ObjectId(str(value))


def _text(value, limit=5000):
    if value is None:
        return ""

    return str(value).strip()[:limit]


def _bool(value):
    if isinstance(value, bool):
        return value

    return str(value or "").strip().lower() in {
        "1",
        "true",
        "yes",
        "on",
    }


def _clean_list(value, limit=100, item_limit=200):
    if value is None:
        return []

    if isinstance(value, str):
        items = [
            item.strip()
            for item in re.split(r"[,\n]", value)
        ]
    elif isinstance(value, (list, tuple)):
        items = value
    else:
        items = []

    output = []
    seen = set()

    for raw in items:
        if isinstance(raw, dict):
            continue

        item = _text(raw, item_limit)
        key = item.lower()

        if not item or key in seen:
            continue

        seen.add(key)
        output.append(item)

        if len(output) >= limit:
            break

    return output


def _clean_bullets(value):
    return _clean_list(
        value,
        limit=30,
        item_limit=700,
    )


def _clean_experience(value):
    if not isinstance(value, list):
        return []

    output = []

    for item in value[:30]:
        if not isinstance(item, dict):
            continue

        output.append(
            {
                "id": _text(
                    item.get("id"),
                    80,
                ),
                "company": _text(
                    item.get("company"),
                    180,
                ),
                "title": _text(
                    item.get("title"),
                    180,
                ),
                "location": _text(
                    item.get("location"),
                    180,
                ),
                "start_date": _text(
                    item.get("start_date"),
                    30,
                ),
                "end_date": _text(
                    item.get("end_date"),
                    30,
                ),
                "current": _bool(
                    item.get("current")
                ),
                "summary": _text(
                    item.get("summary"),
                    3000,
                ),
                "bullets": _clean_bullets(
                    item.get("bullets")
                ),
            }
        )

    return output


def _clean_education(value):
    if not isinstance(value, list):
        return []

    output = []

    for item in value[:30]:
        if not isinstance(item, dict):
            continue

        output.append(
            {
                "id": _text(
                    item.get("id"),
                    80,
                ),
                "school": _text(
                    item.get("school"),
                    220,
                ),
                "degree": _text(
                    item.get("degree"),
                    220,
                ),
                "field": _text(
                    item.get("field"),
                    220,
                ),
                "location": _text(
                    item.get("location"),
                    180,
                ),
                "start_date": _text(
                    item.get("start_date"),
                    30,
                ),
                "end_date": _text(
                    item.get("end_date"),
                    30,
                ),
                "details": _text(
                    item.get("details"),
                    2500,
                ),
            }
        )

    return output


def _clean_projects(value):
    if not isinstance(value, list):
        return []

    output = []

    for item in value[:40]:
        if not isinstance(item, dict):
            continue

        output.append(
            {
                "id": _text(
                    item.get("id"),
                    80,
                ),
                "name": _text(
                    item.get("name"),
                    220,
                ),
                "role": _text(
                    item.get("role"),
                    180,
                ),
                "url": _text(
                    item.get("url"),
                    2000,
                ),
                "description": _text(
                    item.get("description"),
                    3000,
                ),
                "bullets": _clean_bullets(
                    item.get("bullets")
                ),
                "technologies": _clean_list(
                    item.get("technologies"),
                    limit=30,
                    item_limit=100,
                ),
            }
        )

    return output


def _clean_certifications(value):
    if not isinstance(value, list):
        return []

    output = []

    for item in value[:50]:
        if not isinstance(item, dict):
            continue

        output.append(
            {
                "id": _text(
                    item.get("id"),
                    80,
                ),
                "name": _text(
                    item.get("name"),
                    220,
                ),
                "issuer": _text(
                    item.get("issuer"),
                    220,
                ),
                "date": _text(
                    item.get("date"),
                    30,
                ),
                "url": _text(
                    item.get("url"),
                    2000,
                ),
            }
        )

    return output


def _clean_languages(value):
    if not isinstance(value, list):
        return []

    output = []

    for item in value[:30]:
        if not isinstance(item, dict):
            continue

        output.append(
            {
                "id": _text(
                    item.get("id"),
                    80,
                ),
                "language": _text(
                    item.get("language"),
                    120,
                ),
                "level": _text(
                    item.get("level"),
                    120,
                ),
            }
        )

    return output


def _clean_profile(value):
    if not isinstance(value, dict):
        value = {}

    return {
        "full_name": _text(
            value.get("full_name"),
            180,
        ),
        "headline": _text(
            value.get("headline"),
            240,
        ),
        "email": _text(
            value.get("email"),
            320,
        ),
        "phone": _text(
            value.get("phone"),
            80,
        ),
        "location": _text(
            value.get("location"),
            180,
        ),
        "website": _text(
            value.get("website"),
            2000,
        ),
        "linkedin": _text(
            value.get("linkedin"),
            2000,
        ),
        "github": _text(
            value.get("github"),
            2000,
        ),
    }


def _normalize_resume(data, existing=None):
    existing = existing or {}
    updates = {}

    if "name" in data:
        updates["name"] = _text(
            data.get("name"),
            180,
        )

    if "template" in data:
        template = _text(
            data.get("template"),
            40,
        ).lower()

        if template not in RESUME_TEMPLATES:
            raise ValueError(
                "Invalid resume template."
            )

        updates["template"] = template

    if "profile" in data:
        updates["profile"] = _clean_profile(
            data.get("profile")
        )

    if "summary" in data:
        updates["summary"] = _text(
            data.get("summary"),
            6000,
        )

    if "experience" in data:
        updates["experience"] = (
            _clean_experience(
                data.get("experience")
            )
        )

    if "education" in data:
        updates["education"] = (
            _clean_education(
                data.get("education")
            )
        )

    if "skills" in data:
        updates["skills"] = _clean_list(
            data.get("skills"),
            limit=150,
            item_limit=120,
        )

    if "projects" in data:
        updates["projects"] = (
            _clean_projects(
                data.get("projects")
            )
        )

    if "certifications" in data:
        updates["certifications"] = (
            _clean_certifications(
                data.get("certifications")
            )
        )

    if "languages" in data:
        updates["languages"] = (
            _clean_languages(
                data.get("languages")
            )
        )

    if "section_order" in data:
        requested = _clean_list(
            data.get("section_order"),
            limit=20,
            item_limit=40,
        )

        order = [
            key
            for key in requested
            if key in SECTION_KEYS
        ]

        for key in SECTION_KEYS:
            if key not in order:
                order.append(key)

        updates["section_order"] = order

    if "archived" in data:
        updates["archived"] = _bool(
            data.get("archived")
        )

    if "accent" in data:
        updates["accent"] = _text(
            data.get("accent"),
            40,
        )

    name = updates.get(
        "name",
        existing.get("name", ""),
    )

    if not name:
        raise ValueError(
            "Resume name is required."
        )

    return updates


def _resume_defaults(name="Untitled Resume"):
    return {
        "name": name,
        "template": "modern",
        "accent": "primary",
        "profile": {
            "full_name": "",
            "headline": "",
            "email": "",
            "phone": "",
            "location": "",
            "website": "",
            "linkedin": "",
            "github": "",
        },
        "summary": "",
        "experience": [],
        "education": [],
        "skills": [],
        "projects": [],
        "certifications": [],
        "languages": [],
        "section_order": list(
            SECTION_KEYS
        ),
        "archived": False,
    }


def _serialize_resume(doc):
    if not doc:
        return None

    return {
        "id": str(doc["_id"]),
        "name": doc.get(
            "name",
            "Untitled Resume",
        ),
        "template": doc.get(
            "template",
            "modern",
        ),
        "accent": doc.get(
            "accent",
            "primary",
        ),
        "profile": doc.get(
            "profile",
            {},
        ),
        "summary": doc.get(
            "summary",
            "",
        ),
        "experience": doc.get(
            "experience",
            [],
        ),
        "education": doc.get(
            "education",
            [],
        ),
        "skills": doc.get(
            "skills",
            [],
        ),
        "projects": doc.get(
            "projects",
            [],
        ),
        "certifications": doc.get(
            "certifications",
            [],
        ),
        "languages": doc.get(
            "languages",
            [],
        ),
        "section_order": doc.get(
            "section_order",
            list(SECTION_KEYS),
        ),
        "archived": bool(
            doc.get("archived")
        ),
        "created_at": serialize_datetime(
            doc.get("created_at")
        ),
        "updated_at": serialize_datetime(
            doc.get("updated_at")
        ),
    }


def _snapshot_resume(doc):
    return {
        key: copy.deepcopy(value)
        for key, value in doc.items()
        if key
        not in {
            "_id",
            "owner_id",
            "created_at",
            "updated_at",
        }
    }


def _serialize_version(doc):
    return {
        "id": str(doc["_id"]),
        "resume_id": str(
            doc.get("resume_id", "")
        ),
        "version": int(
            doc.get("version", 1) or 1
        ),
        "label": doc.get(
            "label",
            "",
        ),
        "created_at": serialize_datetime(
            doc.get("created_at")
        ),
    }


def _serialize_cover_letter(doc):
    return {
        "id": str(doc["_id"]),
        "name": doc.get(
            "name",
            "Cover Letter",
        ),
        "resume_id": str(
            doc.get("resume_id", "")
        )
        if doc.get("resume_id")
        else "",
        "application_id": str(
            doc.get("application_id", "")
        )
        if doc.get("application_id")
        else "",
        "company": doc.get(
            "company",
            "",
        ),
        "job_title": doc.get(
            "job_title",
            "",
        ),
        "recipient_name": doc.get(
            "recipient_name",
            "",
        ),
        "salutation": doc.get(
            "salutation",
            "Dear Hiring Manager,",
        ),
        "body": doc.get(
            "body",
            "",
        ),
        "closing": doc.get(
            "closing",
            "Sincerely,",
        ),
        "archived": bool(
            doc.get("archived")
        ),
        "created_at": serialize_datetime(
            doc.get("created_at")
        ),
        "updated_at": serialize_datetime(
            doc.get("updated_at")
        ),
    }


def _resume_doc(db, owner_id, resume_id):
    resume_oid = _oid(resume_id)

    if not resume_oid:
        return None

    return db.resumes.find_one(
        {
            "_id": resume_oid,
            "owner_id": owner_id,
        }
    )


def _cover_letter_doc(
    db,
    owner_id,
    cover_letter_id,
):
    cover_oid = _oid(cover_letter_id)

    if not cover_oid:
        return None

    return db.cover_letters.find_one(
        {
            "_id": cover_oid,
            "owner_id": owner_id,
        }
    )


def _ensure_indexes(db):
    try:
        db.resumes.create_index(
            [
                ("owner_id", ASCENDING),
                ("updated_at", DESCENDING),
            ]
        )
        db.resume_versions.create_index(
            [
                ("owner_id", ASCENDING),
                ("resume_id", ASCENDING),
                ("version", DESCENDING),
            ]
        )
        db.cover_letters.create_index(
            [
                ("owner_id", ASCENDING),
                ("updated_at", DESCENDING),
            ]
        )
    except Exception:
        pass


@api_view(["GET", "POST"])
def resumes(request):
    db = get_db()
    owner_id = _owner_oid(request)
    _ensure_indexes(db)

    if request.method == "POST":
        name = _text(
            request.data.get("name"),
            180,
        ) or "Untitled Resume"

        defaults = _resume_defaults(name)
        now = utcnow()

        doc = {
            "owner_id": owner_id,
            **defaults,
            "created_at": now,
            "updated_at": now,
        }

        try:
            updates = _normalize_resume(
                request.data,
                existing=doc,
            )
        except ValueError as exc:
            return Response(
                {"detail": str(exc)},
                status=status.HTTP_400_BAD_REQUEST,
            )

        doc.update(updates)

        result = db.resumes.insert_one(doc)
        doc["_id"] = result.inserted_id

        return Response(
            {
                "resume": _serialize_resume(
                    doc
                )
            },
            status=status.HTTP_201_CREATED,
        )

    archived = str(
        request.query_params.get(
            "archived",
            "false",
        )
    ).lower()

    query = {"owner_id": owner_id}

    if archived == "only":
        query["archived"] = True
    elif archived not in {
        "all",
        "1",
        "true",
    }:
        query["archived"] = {
            "$ne": True
        }

    docs = list(
        db.resumes.find(query).sort(
            "updated_at",
            DESCENDING,
        )
    )

    return Response(
        {
            "resumes": [
                _serialize_resume(doc)
                for doc in docs
            ],
            "total": len(docs),
            "templates": sorted(
                RESUME_TEMPLATES
            ),
        }
    )


@api_view(["GET", "PATCH", "DELETE"])
def resume_detail(request, resume_id):
    db = get_db()
    owner_id = _owner_oid(request)
    doc = _resume_doc(
        db,
        owner_id,
        resume_id,
    )

    if not doc:
        return Response(
            {"detail": "Resume not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    if request.method == "GET":
        return Response(
            {
                "resume": _serialize_resume(
                    doc
                )
            }
        )

    if request.method == "DELETE":
        db.resume_versions.delete_many(
            {
                "owner_id": owner_id,
                "resume_id": doc["_id"],
            }
        )
        db.resumes.delete_one(
            {
                "_id": doc["_id"],
                "owner_id": owner_id,
            }
        )
        db.job_applications.update_many(
            {
                "owner_id": owner_id,
                "resume_id": doc["_id"],
            },
            {
                "$unset": {
                    "resume_id": "",
                }
            },
        )
        return Response(
            status=status.HTTP_204_NO_CONTENT
        )

    try:
        updates = _normalize_resume(
            request.data,
            existing=doc,
        )
    except ValueError as exc:
        return Response(
            {"detail": str(exc)},
            status=status.HTTP_400_BAD_REQUEST,
        )

    updates["updated_at"] = utcnow()

    db.resumes.update_one(
        {
            "_id": doc["_id"],
            "owner_id": owner_id,
        },
        {"$set": updates},
    )

    updated = db.resumes.find_one(
        {
            "_id": doc["_id"],
            "owner_id": owner_id,
        }
    )

    return Response(
        {
            "resume": _serialize_resume(
                updated
            )
        }
    )


@api_view(["POST"])
def resume_duplicate(request, resume_id):
    db = get_db()
    owner_id = _owner_oid(request)
    source = _resume_doc(
        db,
        owner_id,
        resume_id,
    )

    if not source:
        return Response(
            {"detail": "Resume not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    now = utcnow()
    duplicate = {
        "owner_id": owner_id,
        **_snapshot_resume(source),
        "name": (
            _text(
                request.data.get("name"),
                180,
            )
            or f"{source.get('name', 'Resume')} Copy"
        ),
        "archived": False,
        "created_at": now,
        "updated_at": now,
    }

    result = db.resumes.insert_one(
        duplicate
    )
    duplicate["_id"] = result.inserted_id

    return Response(
        {
            "resume": _serialize_resume(
                duplicate
            )
        },
        status=status.HTTP_201_CREATED,
    )


@api_view(["GET", "POST"])
def resume_versions(request, resume_id):
    db = get_db()
    owner_id = _owner_oid(request)
    resume = _resume_doc(
        db,
        owner_id,
        resume_id,
    )

    if not resume:
        return Response(
            {"detail": "Resume not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    if request.method == "POST":
        latest = db.resume_versions.find_one(
            {
                "owner_id": owner_id,
                "resume_id": resume["_id"],
            },
            sort=[("version", DESCENDING)],
        )

        version_number = int(
            latest.get("version", 0)
            if latest
            else 0
        ) + 1

        doc = {
            "owner_id": owner_id,
            "resume_id": resume["_id"],
            "version": version_number,
            "label": (
                _text(
                    request.data.get("label"),
                    180,
                )
                or f"Version {version_number}"
            ),
            "snapshot": _snapshot_resume(
                resume
            ),
            "created_at": utcnow(),
        }

        result = (
            db.resume_versions.insert_one(
                doc
            )
        )
        doc["_id"] = result.inserted_id

        return Response(
            {
                "version": _serialize_version(
                    doc
                )
            },
            status=status.HTTP_201_CREATED,
        )

    docs = list(
        db.resume_versions.find(
            {
                "owner_id": owner_id,
                "resume_id": resume["_id"],
            }
        ).sort("version", DESCENDING)
    )

    return Response(
        {
            "versions": [
                _serialize_version(doc)
                for doc in docs
            ]
        }
    )


@api_view(["POST"])
def resume_restore_version(
    request,
    resume_id,
    version_id,
):
    db = get_db()
    owner_id = _owner_oid(request)
    resume = _resume_doc(
        db,
        owner_id,
        resume_id,
    )
    version_oid = _oid(version_id)

    if not resume or not version_oid:
        return Response(
            {"detail": "Version not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    version = db.resume_versions.find_one(
        {
            "_id": version_oid,
            "owner_id": owner_id,
            "resume_id": resume["_id"],
        }
    )

    if not version:
        return Response(
            {"detail": "Version not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    snapshot = copy.deepcopy(
        version.get("snapshot", {})
    )
    snapshot.pop("archived", None)
    snapshot["updated_at"] = utcnow()

    db.resumes.update_one(
        {
            "_id": resume["_id"],
            "owner_id": owner_id,
        },
        {"$set": snapshot},
    )

    updated = db.resumes.find_one(
        {
            "_id": resume["_id"],
            "owner_id": owner_id,
        }
    )

    return Response(
        {
            "resume": _serialize_resume(
                updated
            )
        }
    )


def _normalize_cover_letter(
    data,
    existing=None,
):
    existing = existing or {}
    updates = {}

    text_fields = {
        "name": 180,
        "company": 180,
        "job_title": 220,
        "recipient_name": 180,
        "salutation": 300,
        "body": 20000,
        "closing": 300,
    }

    for field, limit in text_fields.items():
        if field in data:
            updates[field] = _text(
                data.get(field),
                limit,
            )

    for field in (
        "resume_id",
        "application_id",
    ):
        if field in data:
            value = _oid(
                data.get(field)
            )
            updates[field] = value

    if "archived" in data:
        updates["archived"] = _bool(
            data.get("archived")
        )

    name = updates.get(
        "name",
        existing.get("name", ""),
    )

    if not name:
        raise ValueError(
            "Cover letter name is required."
        )

    return updates


@api_view(["GET", "POST"])
def cover_letters(request):
    db = get_db()
    owner_id = _owner_oid(request)
    _ensure_indexes(db)

    if request.method == "POST":
        now = utcnow()
        doc = {
            "owner_id": owner_id,
            "name": (
                _text(
                    request.data.get("name"),
                    180,
                )
                or "Cover Letter"
            ),
            "resume_id": None,
            "application_id": None,
            "company": "",
            "job_title": "",
            "recipient_name": "",
            "salutation": (
                "Dear Hiring Manager,"
            ),
            "body": "",
            "closing": "Sincerely,",
            "archived": False,
            "created_at": now,
            "updated_at": now,
        }

        try:
            updates = (
                _normalize_cover_letter(
                    request.data,
                    existing=doc,
                )
            )
        except ValueError as exc:
            return Response(
                {"detail": str(exc)},
                status=status.HTTP_400_BAD_REQUEST,
            )

        doc.update(updates)

        result = (
            db.cover_letters.insert_one(
                doc
            )
        )
        doc["_id"] = result.inserted_id

        return Response(
            {
                "cover_letter": (
                    _serialize_cover_letter(
                        doc
                    )
                )
            },
            status=status.HTTP_201_CREATED,
        )

    docs = list(
        db.cover_letters.find(
            {
                "owner_id": owner_id,
                "archived": {"$ne": True},
            }
        ).sort("updated_at", DESCENDING)
    )

    return Response(
        {
            "cover_letters": [
                _serialize_cover_letter(
                    doc
                )
                for doc in docs
            ]
        }
    )


@api_view(["GET", "PATCH", "DELETE"])
def cover_letter_detail(
    request,
    cover_letter_id,
):
    db = get_db()
    owner_id = _owner_oid(request)
    doc = _cover_letter_doc(
        db,
        owner_id,
        cover_letter_id,
    )

    if not doc:
        return Response(
            {
                "detail": (
                    "Cover letter not found."
                )
            },
            status=status.HTTP_404_NOT_FOUND,
        )

    if request.method == "GET":
        return Response(
            {
                "cover_letter": (
                    _serialize_cover_letter(
                        doc
                    )
                )
            }
        )

    if request.method == "DELETE":
        db.cover_letters.delete_one(
            {
                "_id": doc["_id"],
                "owner_id": owner_id,
            }
        )
        db.job_applications.update_many(
            {
                "owner_id": owner_id,
                "cover_letter_id": doc[
                    "_id"
                ],
            },
            {
                "$unset": {
                    "cover_letter_id": "",
                }
            },
        )
        return Response(
            status=status.HTTP_204_NO_CONTENT
        )

    try:
        updates = _normalize_cover_letter(
            request.data,
            existing=doc,
        )
    except ValueError as exc:
        return Response(
            {"detail": str(exc)},
            status=status.HTTP_400_BAD_REQUEST,
        )

    updates["updated_at"] = utcnow()

    db.cover_letters.update_one(
        {
            "_id": doc["_id"],
            "owner_id": owner_id,
        },
        {"$set": updates},
    )

    updated = (
        db.cover_letters.find_one(
            {
                "_id": doc["_id"],
                "owner_id": owner_id,
            }
        )
    )

    return Response(
        {
            "cover_letter": (
                _serialize_cover_letter(
                    updated
                )
            )
        }
    )


def _resume_text(resume):
    parts = []

    profile = resume.get(
        "profile",
        {},
    )

    parts.extend(
        [
            profile.get("headline", ""),
            resume.get("summary", ""),
        ]
    )

    for item in resume.get(
        "experience",
        [],
    ):
        parts.extend(
            [
                item.get("title", ""),
                item.get("company", ""),
                item.get("summary", ""),
                " ".join(
                    item.get("bullets", [])
                ),
            ]
        )

    for item in resume.get(
        "education",
        [],
    ):
        parts.extend(
            [
                item.get("degree", ""),
                item.get("field", ""),
                item.get("school", ""),
                item.get("details", ""),
            ]
        )

    parts.append(
        " ".join(
            resume.get("skills", [])
        )
    )

    for item in resume.get(
        "projects",
        [],
    ):
        parts.extend(
            [
                item.get("name", ""),
                item.get("role", ""),
                item.get("description", ""),
                " ".join(
                    item.get(
                        "technologies",
                        [],
                    )
                ),
                " ".join(
                    item.get("bullets", [])
                ),
            ]
        )

    for item in resume.get(
        "certifications",
        [],
    ):
        parts.extend(
            [
                item.get("name", ""),
                item.get("issuer", ""),
            ]
        )

    for item in resume.get(
        "languages",
        [],
    ):
        parts.extend(
            [
                item.get("language", ""),
                item.get("level", ""),
            ]
        )

    return " ".join(
        part
        for part in parts
        if part
    )


ATS_NOISE_WORDS = {
    "ability", "abilities", "about", "across", "also", "any", "apply",
    "applicant", "applicants", "appropriate", "background", "benefits",
    "best", "business", "collaborate", "communication", "company",
    "considered", "day", "days", "demonstrated", "description",
    "detail", "details", "duties", "environment", "excellent",
    "expected", "familiar", "highly", "ideal", "including", "industry",
    "job", "knowledge", "looking", "minimum", "nice", "opportunity",
    "preferred", "qualification", "qualifications", "required",
    "requirement", "requirements", "responsibilities", "responsibility",
    "role", "strong", "successful", "support", "team", "teams",
    "understanding", "work", "working", "years", "year", "must",
    "proficient", "proficiency", "familiarity", "hands", "hands-on",
}

REQUIREMENT_MARKERS = (
    "required",
    "requirement",
    "requirements",
    "must",
    "minimum",
    "qualification",
    "qualifications",
    "proficient",
    "proficiency",
    "experience with",
    "experience in",
    "knowledge of",
    "familiarity with",
    "hands-on",
    "hands on",
)

KNOWN_SKILL_PHRASES = {
    "machine learning",
    "deep learning",
    "natural language processing",
    "computer vision",
    "data analysis",
    "data science",
    "project management",
    "product management",
    "financial analysis",
    "quality assurance",
    "large language model",
    "large language models",
    "generative ai",
    "artificial intelligence",
    "software engineering",
    "cloud computing",
    "business intelligence",
    "power bi",
    "microsoft excel",
    "google cloud",
    "amazon web services",
    "rest api",
    "rest apis",
    "unit testing",
    "version control",
    "continuous integration",
    "continuous deployment",
    "data engineering",
    "data visualization",
    "feature engineering",
    "statistical analysis",
    "prompt engineering",
    "model evaluation",
    "model training",
    "web development",
    "frontend development",
    "backend development",
    "full stack",
    "full-stack",
}


def _tokens(text):
    return [
        token
        for token in re.findall(
            r"[a-zA-Z][a-zA-Z0-9+#.\-]{1,}",
            str(text or "").lower(),
        )
        if token not in STOP_WORDS
        and token not in ATS_NOISE_WORDS
        and len(token) >= 2
    ]


def _requirement_segments(text):
    segments = re.split(
        r"[\n\r]+|(?<=[.!?;])\s+",
        str(text or ""),
    )

    return [
        segment.strip()
        for segment in segments
        if segment.strip()
        and any(
            marker in segment.lower()
            for marker in REQUIREMENT_MARKERS
        )
    ]


def _weighted_job_terms(job_description, limit=70):
    job_counter = Counter(
        _tokens(job_description)
    )

    required_counter = Counter()

    for segment in _requirement_segments(
        job_description
    ):
        required_counter.update(
            _tokens(segment)
        )

    weighted = []

    for token, frequency in job_counter.items():
        weight = 1.0

        # Repeated terms usually describe the core of the role, but cap the
        # bonus so repeated boilerplate cannot dominate the score.
        weight += min(
            max(frequency - 1, 0),
            3,
        ) * 0.2

        # Requirement language is intentionally weighted more heavily than
        # general descriptive text.
        if required_counter.get(token):
            weight += 0.8
            weight += min(
                required_counter[token] - 1,
                2,
            ) * 0.15

        weighted.append(
            (token, round(weight, 3))
        )

    weighted.sort(
        key=lambda item: (
            -item[1],
            -job_counter[item[0]],
            item[0],
        )
    )

    return weighted[:limit], required_counter


def _weighted_coverage(
    weighted_terms,
    candidate_tokens,
):
    if not weighted_terms:
        return 0.0

    candidate_tokens = set(
        candidate_tokens
    )
    total_weight = sum(
        weight
        for _token, weight
        in weighted_terms
    )

    matched_weight = sum(
        weight
        for token, weight
        in weighted_terms
        if token in candidate_tokens
    )

    return round(
        matched_weight
        / max(total_weight, 1)
        * 100,
        1,
    )


def _resume_evidence_text(resume):
    profile = resume.get(
        "profile",
        {},
    )
    parts = [
        profile.get("headline", ""),
        resume.get("summary", ""),
    ]

    for item in resume.get(
        "experience",
        [],
    ):
        parts.extend(
            [
                item.get("title", ""),
                item.get("summary", ""),
                " ".join(
                    item.get("bullets", [])
                ),
            ]
        )

    for item in resume.get(
        "projects",
        [],
    ):
        parts.extend(
            [
                item.get("name", ""),
                item.get("role", ""),
                item.get("description", ""),
                " ".join(
                    item.get(
                        "technologies",
                        [],
                    )
                ),
                " ".join(
                    item.get("bullets", [])
                ),
            ]
        )

    return " ".join(
        str(part)
        for part in parts
        if part
    )


def _job_phrases(job_description):
    raw = str(
        job_description or ""
    ).lower()
    phrases = {
        phrase
        for phrase in KNOWN_SKILL_PHRASES
        if phrase in raw
    }

    # Pull compact bigrams/trigrams out of requirement sentences. This adds
    # job-specific phrases without relying on a fixed technology dictionary.
    for segment in _requirement_segments(
        job_description
    ):
        words = re.findall(
            r"[a-zA-Z][a-zA-Z0-9+#.\-]{1,}",
            segment.lower(),
        )

        for size in (2, 3):
            for index in range(
                0,
                len(words) - size + 1,
            ):
                window = words[
                    index:index + size
                ]

                if any(
                    word in STOP_WORDS
                    or word in ATS_NOISE_WORDS
                    for word in window
                ):
                    continue

                phrase = " ".join(window)

                if len(phrase) >= 5:
                    phrases.add(phrase)

    return sorted(
        phrases,
        key=lambda phrase: (
            -len(phrase.split()),
            phrase,
        ),
    )[:30]


def _resume_match_completeness(resume):
    profile = resume.get(
        "profile",
        {},
    )
    experience = resume.get(
        "experience",
        [],
    )
    skills = resume.get(
        "skills",
        [],
    )

    experience_bullets = [
        bullet
        for item in experience
        for bullet in item.get(
            "bullets",
            [],
        )
        if str(bullet).strip()
    ]

    checks = {
        "contact": bool(
            profile.get("email")
            and profile.get("full_name")
        ),
        "headline": bool(
            profile.get("headline")
        ),
        "summary": 60
        <= len(
            resume.get(
                "summary",
                "",
            )
        )
        <= 1200,
        "experience": bool(
            experience
        ),
        "experience_detail": (
            len(experience_bullets)
            >= 2
        ),
        "education": bool(
            resume.get("education")
        ),
        "skills": (
            5 <= len(skills) <= 50
        ),
        "results": any(
            re.search(
                r"\b\d+(?:\.\d+)?%|\b\d+[kmb]?\+?\b",
                bullet,
                re.IGNORECASE,
            )
            for bullet in experience_bullets
        ),
    }

    weights = {
        "contact": 8,
        "headline": 10,
        "summary": 14,
        "experience": 18,
        "experience_detail": 18,
        "education": 8,
        "skills": 12,
        "results": 12,
    }

    score = round(
        sum(
            weights[key]
            for key, passed
            in checks.items()
            if passed
        ),
        1,
    )

    return checks, score


def _analyze_resume_match(
    resume,
    job_description,
):
    resume_text = _resume_text(
        resume
    )
    resume_tokens = set(
        _tokens(resume_text)
    )
    evidence_text = (
        _resume_evidence_text(
            resume
        )
    )
    evidence_tokens = set(
        _tokens(evidence_text)
    )

    weighted_terms, required_counter = (
        _weighted_job_terms(
            job_description
        )
    )

    weighted_keywords = [
        token
        for token, _weight
        in weighted_terms
    ]

    matched = [
        token
        for token
        in weighted_keywords
        if token in resume_tokens
    ]
    missing = [
        token
        for token
        in weighted_keywords
        if token not in resume_tokens
    ]

    keyword_score = (
        _weighted_coverage(
            weighted_terms,
            resume_tokens,
        )
    )

    evidence_score = (
        _weighted_coverage(
            weighted_terms,
            evidence_tokens,
        )
    )

    required_terms = [
        (
            token,
            1.0
            + min(
                required_counter[token] - 1,
                2,
            )
            * 0.25,
        )
        for token, _count
        in required_counter.most_common(40)
        if token
        not in ATS_NOISE_WORDS
        and token not in STOP_WORDS
    ]

    required_score = (
        _weighted_coverage(
            required_terms,
            resume_tokens,
        )
        if required_terms
        else keyword_score
    )

    job_phrases = _job_phrases(
        job_description
    )
    resume_lower = (
        resume_text.lower()
    )

    matched_phrases = [
        phrase
        for phrase in job_phrases
        if phrase in resume_lower
    ]
    missing_phrases = [
        phrase
        for phrase in job_phrases
        if phrase not in resume_lower
    ]

    if job_phrases:
        phrase_score = round(
            len(matched_phrases)
            / len(job_phrases)
            * 100,
            1,
        )
    else:
        # When there are no meaningful multi-word requirements, do not hand
        # out a free 100. Use evidence coverage as the conservative fallback.
        phrase_score = evidence_score

    completeness_checks, (
        completeness_score
    ) = _resume_match_completeness(
        resume
    )

    overall = round(
        keyword_score * 0.32
        + required_score * 0.25
        + evidence_score * 0.23
        + phrase_score * 0.12
        + completeness_score * 0.08,
        1,
    )

    # A resume that only lists keywords in a Skills section should not score
    # like one that demonstrates those requirements in experience/projects.
    evidence_gap = (
        keyword_score
        - evidence_score
    )

    stuffing_penalty = 0.0

    if (
        keyword_score >= 35
        and evidence_gap >= 18
    ):
        stuffing_penalty = min(
            15.0,
            round(
                (evidence_gap - 12)
                * 0.35,
                1,
            ),
        )
        overall = round(
            max(
                0.0,
                overall
                - stuffing_penalty,
            ),
            1,
        )

    # 100 should mean essentially every relevant requirement is covered and
    # evidenced. Do not round a merely strong resume up to a perfect result.
    if overall > 99.0 and not (
        keyword_score == 100.0
        and required_score == 100.0
        and evidence_score == 100.0
        and phrase_score == 100.0
        and completeness_score == 100.0
    ):
        overall = 99.0

    if overall >= 85:
        score_band = "Excellent match"
    elif overall >= 70:
        score_band = "Strong match"
    elif overall >= 50:
        score_band = "Moderate match"
    elif overall >= 30:
        score_band = "Weak match"
    else:
        score_band = "Low match"

    recommendations = []

    if required_score < 70:
        recommendations.append(
            "Several requirement-focused terms from the job description are not covered. Review the missing terms and add only skills or experience you genuinely have."
        )

    if evidence_score + 15 < keyword_score:
        recommendations.append(
            "Some matched terms appear mainly in skills or profile text. Strengthen relevant experience/project bullets with truthful evidence of using those skills."
        )

    if phrase_score < 60 and missing_phrases:
        recommendations.append(
            (
                "Review missing job-specific phrases: "
                + ", ".join(
                    missing_phrases[:8]
                )
            )
        )

    if not completeness_checks[
        "summary"
    ]:
        recommendations.append(
            "Add a concise professional summary tailored to the target role."
        )

    if not completeness_checks[
        "experience_detail"
    ]:
        recommendations.append(
            "Add concrete experience bullets that show what you did and the outcome."
        )

    if not completeness_checks[
        "results"
    ]:
        recommendations.append(
            "Where accurate, add measurable outcomes to experience bullets."
        )

    if missing:
        recommendations.append(
            (
                "Review missing high-value keywords for truthful coverage: "
                + ", ".join(
                    missing[:15]
                )
            )
        )

    critical_missing = [
        token
        for token, _weight
        in weighted_terms
        if token in missing
    ][:12]

    return {
        "score": overall,
        "score_band": score_band,
        "keyword_score": keyword_score,
        "required_score": required_score,
        "evidence_score": evidence_score,
        "phrase_score": phrase_score,
        "completeness_score": (
            completeness_score
        ),
        "matched_keywords": matched[:40],
        "missing_keywords": missing[:40],
        "critical_missing": (
            critical_missing
        ),
        "matched_phrases": (
            matched_phrases
        ),
        "missing_phrases": (
            missing_phrases
        ),
        "completeness": (
            completeness_checks
        ),
        "recommendations": (
            recommendations
        ),
        "diagnostics": {
            "job_terms_analyzed": len(
                weighted_terms
            ),
            "required_terms_analyzed": len(
                required_terms
            ),
            "phrases_analyzed": len(
                job_phrases
            ),
            "matched_job_terms": len(
                matched
            ),
            "evidence_gap": round(
                evidence_gap,
                1,
            ),
            "keyword_stuffing_penalty": (
                stuffing_penalty
            ),
        },
        "disclaimer": (
            "Revnivo calculates a deterministic resume-to-job relevance score from requirement, keyword, phrase, and evidence coverage. It does not reproduce a proprietary employer ATS or predict hiring outcomes."
        ),
    }


@api_view(["POST"])
def resume_analyze(request, resume_id):
    db = get_db()
    owner_id = _owner_oid(request)
    resume = _resume_doc(
        db,
        owner_id,
        resume_id,
    )

    if not resume:
        return Response(
            {"detail": "Resume not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    job_description = _text(
        request.data.get(
            "job_description"
        ),
        40000,
    )

    meaningful_tokens = _tokens(
        job_description
    )

    if (
        len(job_description) < 100
        or len(
            set(
                meaningful_tokens
            )
        ) < 15
    ):
        return Response(
            {
                "detail": (
                    "Paste a fuller job description with at least "
                    "15 meaningful terms so the match score is reliable."
                )
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    return Response(
        _analyze_resume_match(
            resume,
            job_description,
        )
    )


@api_view(["PATCH"])
def application_materials(
    request,
    application_id,
):
    db = get_db()
    owner_id = _owner_oid(request)
    application_oid = _oid(
        application_id
    )

    if not application_oid:
        return Response(
            {
                "detail": (
                    "Application not found."
                )
            },
            status=status.HTTP_404_NOT_FOUND,
        )

    application = (
        db.job_applications.find_one(
            {
                "_id": application_oid,
                "owner_id": owner_id,
            }
        )
    )

    if not application:
        return Response(
            {
                "detail": (
                    "Application not found."
                )
            },
            status=status.HTTP_404_NOT_FOUND,
        )

    updates = {}

    if "resume_id" in request.data:
        raw_resume_id = (
            request.data.get("resume_id")
        )

        if raw_resume_id:
            resume = _resume_doc(
                db,
                owner_id,
                raw_resume_id,
            )

            if not resume:
                return Response(
                    {
                        "detail": (
                            "Selected resume not found."
                        )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            updates["resume_id"] = resume[
                "_id"
            ]
        else:
            updates["resume_id"] = None

    if "cover_letter_id" in request.data:
        raw_cover_id = request.data.get(
            "cover_letter_id"
        )

        if raw_cover_id:
            cover = _cover_letter_doc(
                db,
                owner_id,
                raw_cover_id,
            )

            if not cover:
                return Response(
                    {
                        "detail": (
                            "Selected cover letter not found."
                        )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            updates[
                "cover_letter_id"
            ] = cover["_id"]
        else:
            updates[
                "cover_letter_id"
            ] = None

    updates["updated_at"] = utcnow()

    db.job_applications.update_one(
        {
            "_id": application_oid,
            "owner_id": owner_id,
        },
        {"$set": updates},
    )

    return Response(
        {
            "resume_id": str(
                updates.get(
                    "resume_id",
                    application.get(
                        "resume_id",
                        "",
                    ),
                )
                or ""
            ),
            "cover_letter_id": str(
                updates.get(
                    "cover_letter_id",
                    application.get(
                        "cover_letter_id",
                        "",
                    ),
                )
                or ""
            ),
        }
    )


@api_view(["GET"])
def resume_job_matches(request, resume_id):
    from . import jobs as jobs_module

    db = get_db()
    owner_id = _owner_oid(request)
    resume = _resume_doc(
        db,
        owner_id,
        resume_id,
    )

    if not resume:
        return Response(
            {"detail": "Resume not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    with jobs_module._cache_lock:
        payload = copy.deepcopy(
            jobs_module._cache.get("payload")
        )

    if not payload:
        return Response(
            {
                "matches": [],
                "total_jobs": 0,
                "detail": (
                    "Refresh the Jobs board first so "
                    "Revnivo has current listings to compare."
                ),
            }
        )

    resume_text = _resume_text(resume)
    resume_tokens = set(
        _tokens(resume_text)
    )
    resume_skills = {
        skill.lower()
        for skill in resume.get(
            "skills",
            [],
        )
        if skill
    }

    matches = []

    for job in payload.get("jobs", []):
        title = str(
            job.get("title") or ""
        )
        category = str(
            job.get("category") or ""
        )
        summary = str(
            job.get("summary") or ""
        )
        employment = str(
            job.get(
                "employment_type"
            )
            or ""
        )

        job_text = " ".join(
            [
                title,
                category,
                summary,
                employment,
            ]
        )

        job_tokens = set(
            _tokens(job_text)
        )

        if not job_tokens:
            continue

        overlap = sorted(
            resume_tokens & job_tokens
        )

        keyword_coverage = (
            len(overlap)
            / max(
                min(
                    len(job_tokens),
                    60,
                ),
                1,
            )
        )

        title_tokens = set(
            _tokens(title)
        )
        title_overlap = (
            len(
                title_tokens
                & resume_tokens
            )
            / max(
                len(title_tokens),
                1,
            )
        )

        category_tokens = set(
            _tokens(category)
        )
        category_overlap = (
            len(
                category_tokens
                & resume_tokens
            )
            / max(
                len(category_tokens),
                1,
            )
            if category_tokens
            else 0
        )

        skill_matches = [
            skill
            for skill in resume_skills
            if skill
            and skill in job_text.lower()
        ]

        skill_bonus = min(
            len(skill_matches) / 10,
            1,
        )

        score = round(
            min(
                100,
                (
                    keyword_coverage
                    * 45
                    + title_overlap
                    * 30
                    + category_overlap
                    * 15
                    + skill_bonus
                    * 10
                ),
            ),
            1,
        )

        if score <= 0:
            continue

        matches.append(
            {
                "job": job,
                "score": score,
                "matched_keywords": overlap[
                    :20
                ],
                "matched_skills": (
                    skill_matches[:15]
                ),
            }
        )

    matches.sort(
        key=lambda item: (
            -item["score"],
            str(
                item["job"].get(
                    "platform",
                    "",
                )
            ).lower(),
            str(
                item["job"].get(
                    "title",
                    "",
                )
            ).lower(),
        )
    )

    limit = request.query_params.get(
        "limit",
        "100",
    )

    try:
        limit = int(limit)
    except (TypeError, ValueError):
        limit = 100

    limit = max(
        1,
        min(limit, 250),
    )

    return Response(
        {
            "matches": matches[:limit],
            "total_jobs": len(
                payload.get(
                    "jobs",
                    [],
                )
            ),
            "resume_id": str(
                resume["_id"]
            ),
            "resume_name": resume.get(
                "name",
                "",
            ),
            "method": (
                "Heuristic relevance based on truthful "
                "resume/job keyword and skill overlap."
            ),
        }
    )


IMPORT_SECTION_ALIASES = {
    "summary": {
        "summary",
        "profile",
        "professional summary",
        "professional profile",
        "career summary",
        "about",
        "about me",
        "objective",
        "career objective",
    },
    "experience": {
        "experience",
        "work experience",
        "professional experience",
        "employment",
        "employment history",
        "career history",
        "work history",
    },
    "education": {
        "education",
        "academic background",
        "academic qualifications",
        "qualifications",
    },
    "skills": {
        "skills",
        "technical skills",
        "core skills",
        "key skills",
        "competencies",
        "core competencies",
        "technologies",
        "tools",
    },
    "projects": {
        "projects",
        "selected projects",
        "personal projects",
        "academic projects",
    },
    "certifications": {
        "certifications",
        "certificates",
        "licenses",
        "licenses & certifications",
        "courses",
        "training",
    },
    "languages": {
        "languages",
        "language",
        "language skills",
    },
}

IMPORT_HEADING_TO_SECTION = {
    alias: key
    for key, aliases in IMPORT_SECTION_ALIASES.items()
    for alias in aliases
}


def _import_heading_key(line):
    cleaned = re.sub(
        r"[^a-zA-Z0-9& ]+",
        " ",
        str(line or "").strip().lower(),
    )
    cleaned = re.sub(r"\s+", " ", cleaned).strip()

    if cleaned in IMPORT_HEADING_TO_SECTION:
        return IMPORT_HEADING_TO_SECTION[cleaned]

    return ""


def _extract_resume_upload_text(uploaded):
    file_name = str(
        getattr(uploaded, "name", "resume")
        or "resume"
    )
    content_type = str(
        getattr(uploaded, "content_type", "")
        or ""
    ).split(";", 1)[0].lower()

    data = uploaded.read()
    max_bytes = 12 * 1024 * 1024

    if not data:
        raise ValueError("The uploaded CV is empty.")

    if len(data) > max_bytes:
        raise ValueError(
            "CV file is too large. Maximum size is 12 MB."
        )

    lower_name = file_name.lower()

    if (
        content_type == "application/pdf"
        or lower_name.endswith(".pdf")
    ):
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError(
                "PDF import dependency is not installed. "
                "Run pip install -r requirements.txt."
            ) from exc

        try:
            reader = PdfReader(BytesIO(data))
            pages = []

            for page in reader.pages[:30]:
                text = page.extract_text() or ""
                if text.strip():
                    pages.append(text)

            extracted = "\n\n".join(pages)
        except Exception as exc:
            raise ValueError(
                "Could not read this PDF CV."
            ) from exc

    elif (
        content_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or lower_name.endswith(".docx")
    ):
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError(
                "DOCX import dependency is not installed. "
                "Run pip install -r requirements.txt."
            ) from exc

        try:
            document = Document(BytesIO(data))
            chunks = []

            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if text:
                    chunks.append(text)
                else:
                    chunks.append("")

            for table in document.tables:
                for row in table.rows:
                    values = [
                        cell.text.strip()
                        for cell in row.cells
                        if cell.text.strip()
                    ]
                    if values:
                        chunks.append(" | ".join(values))

            extracted = "\n".join(chunks)
        except Exception as exc:
            raise ValueError(
                "Could not read this DOCX CV."
            ) from exc

    elif (
        content_type in {
            "text/plain",
            "text/markdown",
        }
        or lower_name.endswith(".txt")
        or lower_name.endswith(".md")
    ):
        try:
            extracted = data.decode("utf-8")
        except UnicodeDecodeError:
            try:
                extracted = data.decode(
                    "utf-8-sig"
                )
            except UnicodeDecodeError as exc:
                raise ValueError(
                    "TXT CV must use UTF-8 encoding."
                ) from exc

    else:
        raise ValueError(
            "Unsupported CV format. Upload PDF, DOCX, or TXT."
        )

    extracted = (
        extracted.replace("\r\n", "\n")
        .replace("\r", "\n")
        .replace("\u00a0", " ")
    )
    extracted = re.sub(
        r"[ \t]+",
        " ",
        extracted,
    )
    extracted = re.sub(
        r"\n{4,}",
        "\n\n\n",
        extracted,
    ).strip()

    if len(extracted) < 40:
        raise ValueError(
            "Very little text could be extracted from this CV. "
            "If it is a scanned image PDF, export it as a text-based PDF or DOCX."
        )

    return extracted, file_name


def _split_import_sections(text):
    sections = {
        "header": [],
        "summary": [],
        "experience": [],
        "education": [],
        "skills": [],
        "projects": [],
        "certifications": [],
        "languages": [],
    }

    current = "header"

    for raw_line in text.split("\n"):
        line = raw_line.strip()

        if not line:
            sections[current].append("")
            continue

        heading = _import_heading_key(line)

        if heading:
            current = heading
            continue

        sections[current].append(line)

    return sections


def _first_email(text):
    match = re.search(
        r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b",
        text,
        flags=re.IGNORECASE,
    )
    return match.group(0) if match else ""


def _first_phone(text):
    matches = re.findall(
        r"(?<!\w)(?:\+?\d[\d\s().-]{7,}\d)",
        text,
    )

    for match in matches:
        digits = re.sub(
            r"\D",
            "",
            match,
        )
        if 8 <= len(digits) <= 16:
            return re.sub(
                r"\s+",
                " ",
                match,
            ).strip()

    return ""


def _first_url(text, host_hint=""):
    pattern = (
        r"(?:https?://)?(?:www\.)?"
        r"[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}"
        r"(?:/[^\s,;]*)?"
    )

    for match in re.findall(
        pattern,
        text,
        flags=re.IGNORECASE,
    ):
        if host_hint and host_hint not in match.lower():
            continue

        return match.rstrip(").,;")

    return ""


def _profile_from_import(sections, full_text):
    header_lines = [
        line
        for line in sections.get(
            "header",
            [],
        )
        if line
    ][:14]

    email = _first_email(full_text)
    phone = _first_phone(full_text)
    linkedin = _first_url(
        full_text,
        "linkedin.",
    )
    github = _first_url(
        full_text,
        "github.",
    )

    contact_fragments = {
        email.lower(),
        phone.lower(),
        linkedin.lower(),
        github.lower(),
    }

    candidates = []

    for line in header_lines:
        lower = line.lower()

        if (
            "@" in line
            or "linkedin." in lower
            or "github." in lower
            or re.search(
                r"\+?\d[\d\s().-]{7,}",
                line,
            )
        ):
            continue

        if _import_heading_key(line):
            continue

        if (
            len(line) < 2
            or len(line) > 180
        ):
            continue

        if lower in contact_fragments:
            continue

        candidates.append(line)

    full_name = (
        candidates[0]
        if candidates
        else ""
    )

    headline = (
        candidates[1]
        if len(candidates) > 1
        else ""
    )

    website = ""

    for url in re.findall(
        r"(?:https?://)?(?:www\.)?"
        r"[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}"
        r"(?:/[^\s,;]*)?",
        full_text,
        flags=re.IGNORECASE,
    ):
        lower = url.lower()
        if (
            "linkedin." not in lower
            and "github." not in lower
            and "@" not in lower
        ):
            website = url.rstrip(").,;")
            break

    location = ""

    for line in header_lines:
        if line in {
            full_name,
            headline,
        }:
            continue

        if (
            email
            and email.lower()
            in line.lower()
        ):
            continue

        if (
            phone
            and re.sub(r"\D", "", phone)
            in re.sub(r"\D", "", line)
        ):
            continue

        lower = line.lower()

        if (
            "linkedin" in lower
            or "github" in lower
            or "http" in lower
            or "www." in lower
        ):
            continue

        if (
            "," in line
            and len(line) <= 100
        ):
            location = line
            break

    return {
        "full_name": full_name,
        "headline": headline,
        "email": email,
        "phone": phone,
        "location": location,
        "website": website,
        "linkedin": linkedin,
        "github": github,
    }


def _compact_section_text(lines):
    chunks = []
    paragraph = []

    for line in lines:
        if not line:
            if paragraph:
                chunks.append(
                    " ".join(paragraph)
                )
                paragraph = []
            continue

        paragraph.append(line)

    if paragraph:
        chunks.append(
            " ".join(paragraph)
        )

    return "\n".join(chunks).strip()


def _parse_import_skills(lines):
    raw = " ".join(
        line
        for line in lines
        if line
    )
    raw = raw.replace("•", ",")
    raw = raw.replace("|", ",")
    raw = raw.replace(";", ",")

    values = []

    for item in re.split(
        r",|\s{2,}",
        raw,
    ):
        item = item.strip(" -–—•\t")
        if (
            item
            and len(item) <= 120
        ):
            values.append(item)

    return _clean_list(
        values,
        limit=100,
        item_limit=120,
    )


def _blocks_from_lines(lines):
    blocks = []
    current = []

    for line in lines:
        if not line:
            if current:
                blocks.append(current)
                current = []
            continue

        current.append(line)

    if current:
        blocks.append(current)

    if (
        len(blocks) <= 1
        and len(
            [
                line
                for line in lines
                if line
            ]
        ) >= 8
    ):
        non_empty = [
            line
            for line in lines
            if line
        ]
        blocks = [
            non_empty[index:index + 5]
            for index in range(
                0,
                len(non_empty),
                5,
            )
        ]

    return blocks


def _import_bullets(lines):
    bullets = []

    for line in lines:
        stripped = line.strip()

        if re.match(
            r"^[•·▪◦*-]\s+",
            stripped,
        ):
            value = re.sub(
                r"^[•·▪◦*-]\s+",
                "",
                stripped,
            ).strip()
            if value:
                bullets.append(value)

    return bullets


def _parse_import_experience(lines):
    blocks = _blocks_from_lines(lines)
    output = []

    for index, block in enumerate(
        blocks[:20]
    ):
        clean = [
            line
            for line in block
            if line
        ]

        if not clean:
            continue

        bullets = _import_bullets(clean)
        non_bullets = [
            line
            for line in clean
            if line
            and line not in bullets
            and not re.match(
                r"^[•·▪◦*-]\s+",
                line,
            )
        ]

        title = (
            non_bullets[0]
            if non_bullets
            else "Imported role"
        )
        company = (
            non_bullets[1]
            if len(non_bullets) > 1
            else ""
        )

        date_line = next(
            (
                line
                for line in non_bullets
                if re.search(
                    r"\b(?:19|20)\d{2}\b",
                    line,
                )
            ),
            "",
        )

        summary_lines = [
            line
            for line in non_bullets[2:]
            if line != date_line
        ]

        output.append(
            {
                "id": (
                    f"imported-exp-{index + 1}"
                ),
                "company": company,
                "title": title,
                "location": "",
                "start_date": date_line,
                "end_date": "",
                "current": (
                    "present"
                    in date_line.lower()
                ),
                "summary": " ".join(
                    summary_lines
                ),
                "bullets": bullets,
            }
        )

    return output


def _parse_import_education(lines):
    blocks = _blocks_from_lines(lines)
    output = []

    for index, block in enumerate(
        blocks[:15]
    ):
        clean = [
            line
            for line in block
            if line
        ]

        if not clean:
            continue

        output.append(
            {
                "id": (
                    f"imported-edu-{index + 1}"
                ),
                "school": clean[0],
                "degree": (
                    clean[1]
                    if len(clean) > 1
                    else ""
                ),
                "field": "",
                "location": "",
                "start_date": "",
                "end_date": "",
                "details": " ".join(
                    clean[2:]
                ),
            }
        )

    return output


def _parse_import_projects(lines):
    blocks = _blocks_from_lines(lines)
    output = []

    for index, block in enumerate(
        blocks[:20]
    ):
        clean = [
            line
            for line in block
            if line
        ]

        if not clean:
            continue

        output.append(
            {
                "id": (
                    f"imported-project-{index + 1}"
                ),
                "name": clean[0],
                "role": "",
                "url": next(
                    (
                        item
                        for item in clean
                        if re.search(
                            r"(?:https?://|www\.)",
                            item,
                            re.IGNORECASE,
                        )
                    ),
                    "",
                ),
                "description": " ".join(
                    clean[1:]
                ),
                "bullets": _import_bullets(
                    clean
                ),
                "technologies": [],
            }
        )

    return output


def _parse_import_certifications(lines):
    output = []

    for index, line in enumerate(
        [
            item
            for item in lines
            if item
        ][:40]
    ):
        output.append(
            {
                "id": (
                    f"imported-cert-{index + 1}"
                ),
                "name": line,
                "issuer": "",
                "date": "",
                "url": "",
            }
        )

    return output


def _parse_import_languages(lines):
    output = []

    for index, line in enumerate(
        [
            item
            for item in lines
            if item
        ][:30]
    ):
        parts = re.split(
            r"\s*[-–—:|]\s*",
            line,
            maxsplit=1,
        )

        output.append(
            {
                "id": (
                    f"imported-lang-{index + 1}"
                ),
                "language": parts[0],
                "level": (
                    parts[1]
                    if len(parts) > 1
                    else ""
                ),
            }
        )

    return output


def _resume_from_imported_text(
    text,
    file_name,
):
    sections = _split_import_sections(
        text
    )
    profile = _profile_from_import(
        sections,
        text,
    )

    summary = _compact_section_text(
        sections.get(
            "summary",
            [],
        )
    )

    if not summary:
        header_candidates = [
            line
            for line in sections.get(
                "header",
                [],
            )
            if line
            and line
            not in {
                profile.get("full_name"),
                profile.get("headline"),
            }
        ]

        descriptive = [
            line
            for line in header_candidates
            if len(line) >= 70
        ]

        if descriptive:
            summary = descriptive[0]

    stem = re.sub(
        r"\.(pdf|docx|txt|md)$",
        "",
        file_name,
        flags=re.IGNORECASE,
    ).strip()

    resume = _resume_defaults(
        stem or "Imported Resume"
    )
    resume.update(
        {
            "profile": profile,
            "summary": summary,
            "experience": (
                _parse_import_experience(
                    sections.get(
                        "experience",
                        [],
                    )
                )
            ),
            "education": (
                _parse_import_education(
                    sections.get(
                        "education",
                        [],
                    )
                )
            ),
            "skills": _parse_import_skills(
                sections.get(
                    "skills",
                    [],
                )
            ),
            "projects": (
                _parse_import_projects(
                    sections.get(
                        "projects",
                        [],
                    )
                )
            ),
            "certifications": (
                _parse_import_certifications(
                    sections.get(
                        "certifications",
                        [],
                    )
                )
            ),
            "languages": (
                _parse_import_languages(
                    sections.get(
                        "languages",
                        [],
                    )
                )
            ),
        }
    )

    return resume, sections


def _ats_readiness_report(
    resume,
    raw_text="",
):
    profile = resume.get(
        "profile",
        {},
    )
    experience = resume.get(
        "experience",
        [],
    )
    education = resume.get(
        "education",
        [],
    )
    skills = resume.get(
        "skills",
        [],
    )
    projects = resume.get(
        "projects",
        [],
    )
    certifications = resume.get(
        "certifications",
        [],
    )

    email = str(
        profile.get("email") or ""
    ).strip()
    phone = str(
        profile.get("phone") or ""
    ).strip()
    phone_digits = re.sub(
        r"\D",
        "",
        phone,
    )

    experience_bullets = [
        str(bullet).strip()
        for item in experience
        for bullet in item.get(
            "bullets",
            [],
        )
        if str(bullet).strip()
    ]

    experience_text = " ".join(
        [
            str(
                item.get(
                    "summary",
                    "",
                )
            )
            + " "
            + " ".join(
                item.get(
                    "bullets",
                    [],
                )
            )
            for item in experience
        ]
    )

    quantified = bool(
        re.search(
            r"\b\d+(?:\.\d+)?%|\b\d+[kmb]?\+?\b",
            experience_text,
            flags=re.IGNORECASE,
        )
    )

    raw_word_count = len(
        re.findall(
            r"\b\w+\b",
            str(raw_text or ""),
        )
    )

    structured_experience = any(
        str(
            item.get("title") or ""
        ).strip()
        and str(
            item.get("company") or ""
        ).strip()
        for item in experience
    )

    structured_education = any(
        str(
            item.get("school") or ""
        ).strip()
        and (
            str(
                item.get("degree") or ""
            ).strip()
            or str(
                item.get("field") or ""
            ).strip()
        )
        for item in education
    )

    populated_sections = sum(
        [
            bool(experience),
            bool(education),
            bool(skills),
            bool(projects),
            bool(certifications),
        ]
    )

    checks = {
        "name": bool(
            str(
                profile.get(
                    "full_name"
                )
                or ""
            ).strip()
        ),
        "email": bool(
            re.fullmatch(
                r"[^\s@]+@[^\s@]+\.[^\s@]+",
                email,
            )
        ),
        "phone": (
            8 <= len(phone_digits) <= 16
        ),
        "headline": (
            8
            <= len(
                str(
                    profile.get(
                        "headline"
                    )
                    or ""
                ).strip()
            )
            <= 180
        ),
        "summary": (
            60
            <= len(
                str(
                    resume.get(
                        "summary"
                    )
                    or ""
                ).strip()
            )
            <= 1200
        ),
        "experience": (
            structured_experience
        ),
        "experience_detail": (
            len(
                experience_bullets
            )
            >= 3
        ),
        "education": (
            structured_education
        ),
        "skills": (
            5 <= len(skills) <= 40
        ),
        "quantified_results": (
            quantified
        ),
        "parseable_text": (
            180
            <= raw_word_count
            <= 1800
        ),
        "section_balance": (
            populated_sections >= 3
        ),
    }

    weights = {
        "name": 4,
        "email": 7,
        "phone": 4,
        "headline": 7,
        "summary": 10,
        "experience": 14,
        "experience_detail": 14,
        "education": 7,
        "skills": 10,
        "quantified_results": 8,
        "parseable_text": 8,
        "section_balance": 7,
    }

    raw_score = sum(
        weights[key]
        for key, passed
        in checks.items()
        if passed
    )

    # This is intentionally not presented as an employer ATS score. Without a
    # target job description, the highest meaningful result is general resume
    # readiness, so cap it below 100 to prevent a misleading "perfect ATS" UI.
    score = min(
        95,
        raw_score,
    )

    warnings = []

    if not checks["email"]:
        warnings.append(
            "Add a valid email address."
        )

    if not checks["phone"]:
        warnings.append(
            "Add a valid phone number."
        )

    if not checks["headline"]:
        warnings.append(
            "Add a clear professional headline."
        )

    if not checks["summary"]:
        warnings.append(
            "Use a concise professional summary of roughly 60–1200 characters."
        )

    if not checks["experience"]:
        warnings.append(
            "Review Experience: at least one role should include both a job title and company."
        )

    if not checks[
        "experience_detail"
    ]:
        warnings.append(
            "Add at least three concrete experience bullets so the resume contains enough evidence."
        )

    if not checks["education"]:
        warnings.append(
            "Review Education: include a school plus a degree or field where applicable."
        )

    if not checks["skills"]:
        warnings.append(
            "Keep the skills section focused: about 5–40 relevant skills."
        )

    if not checks[
        "quantified_results"
    ]:
        warnings.append(
            "Where accurate, add measurable outcomes to experience bullets."
        )

    if not checks[
        "parseable_text"
    ]:
        warnings.append(
            "The resume text looks unusually short or long; review parsing and overall length."
        )

    if not checks[
        "section_balance"
    ]:
        warnings.append(
            "Add enough structured sections for a balanced resume."
        )

    return {
        "score": score,
        "raw_score": raw_score,
        "max_general_score": 95,
        "kind": "resume_readiness",
        "checks": checks,
        "warnings": warnings,
        "word_count": raw_word_count,
        "disclaimer": (
            "Resume Readiness measures general structure and parseability only. "
            "It is capped at 95 because no target job description is involved. "
            "Use Job-specific ATS Match for resume-to-role relevance."
        ),
    }


@api_view(["POST"])
@parser_classes(
    [
        MultiPartParser,
        FormParser,
    ]
)
def resume_import(request):
    db = get_db()
    owner_id = _owner_oid(request)
    _ensure_indexes(db)

    uploaded = request.FILES.get(
        "file"
    )

    if not uploaded:
        return Response(
            {
                "detail": (
                    "Choose a PDF, DOCX, or TXT CV to upload."
                )
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    try:
        text, file_name = (
            _extract_resume_upload_text(
                uploaded
            )
        )
        parsed_resume, sections = (
            _resume_from_imported_text(
                text,
                file_name,
            )
        )
    except ValueError as exc:
        return Response(
            {"detail": str(exc)},
            status=status.HTTP_400_BAD_REQUEST,
        )

    requested_name = _text(
        request.data.get("name"),
        180,
    )

    if requested_name:
        parsed_resume[
            "name"
        ] = requested_name

    now = utcnow()
    doc = {
        "owner_id": owner_id,
        **parsed_resume,
        "import_source_name": file_name,
        "created_at": now,
        "updated_at": now,
    }

    result = db.resumes.insert_one(doc)
    doc["_id"] = result.inserted_id

    report = _ats_readiness_report(
        doc,
        raw_text=text,
    )

    detected_sections = [
        key
        for key in SECTION_KEYS
        if sections.get(key)
    ]

    return Response(
        {
            "resume": _serialize_resume(
                doc
            ),
            "import": {
                "file_name": file_name,
                "characters_extracted": len(
                    text
                ),
                "words_extracted": len(
                    re.findall(
                        r"\S+",
                        text,
                    )
                ),
                "detected_sections": (
                    detected_sections
                ),
                "readiness": report,
            },
        },
        status=status.HTTP_201_CREATED,
    )


@api_view(["GET"])
def resume_readiness(request, resume_id):
    db = get_db()
    owner_id = _owner_oid(request)
    resume = _resume_doc(
        db,
        owner_id,
        resume_id,
    )

    if not resume:
        return Response(
            {"detail": "Resume not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    return Response(
        {
            "readiness": _ats_readiness_report(
                resume,
                raw_text=_resume_text(
                    resume
                ),
            )
        }
    )
